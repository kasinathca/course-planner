"""Bounded, non-executing interchange adapters for Phase 6."""

from __future__ import annotations

import ast
import csv
import io
import json
import re
import zipfile
from collections import OrderedDict
from dataclasses import dataclass
from datetime import UTC, datetime
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from ffcs_studio.persistence.database import Database

CSV_COLUMNS = (
    "semester_label",
    "category",
    "course_code",
    "course_title",
    "source_row",
    "slot",
    "venue",
    "faculty",
    "course_type",
)
CATEGORY_FILENAMES = {
    "PC": "ProgrammeCore.ts",
    "PE": "ProgrammeElective.ts",
    "UC": "UniversityCore.ts",
    "UE": "UniversityElective.ts",
    "SPE": "SpecializationElective.ts",
    "BC": "BridgeCourse.ts",
    "NC": "NonCredit.ts",
}
FILENAME_CATEGORIES = {value.casefold(): key for key, value in CATEGORY_FILENAMES.items()}
KNOWN_COMPONENTS = {"TH", "ETH", "ELA", "EPJ", "SS"}


class InterchangeError(ValueError):
    """A bounded interchange document was malformed or unsupported."""


@dataclass(frozen=True, slots=True)
class SemesterRows:
    campus: str
    semester_label: str
    captured_at: str
    courses: tuple[dict[str, Any], ...]

    @property
    def allocation_count(self) -> int:
        return sum(len(course["allocations"]) for course in self.courses)


class SemesterExportService:
    """Export only active, semester-scoped canonical allocation evidence."""

    def __init__(self, database: Database) -> None:
        self.database = database

    def snapshot(self, *, campus: str, semester_label: str) -> SemesterRows:
        with self.database.connection() as connection:
            rows = connection.execute(
                "SELECT o.category, r.raw_course_code, r.raw_course_title, "
                "s.captured_at, a.source_row, a.raw_slot, a.raw_venue, "
                "a.raw_faculty, a.raw_course_type, a.allocation_ordinal "
                "FROM offerings o "
                "JOIN offering_revisions r ON r.revision_id = o.active_revision_id "
                "JOIN import_snapshots s ON s.import_id = r.import_id "
                "LEFT JOIN allocation_rows a ON a.revision_id = r.revision_id "
                "WHERE o.campus = ? AND o.semester_label = ? "
                "ORDER BY o.category, r.raw_course_code, r.course_ordinal, "
                "a.allocation_ordinal",
                (campus, semester_label),
            ).fetchall()
        if not rows:
            raise InterchangeError("No active offerings match that campus and semester")
        grouped: OrderedDict[tuple[str, str, str], dict[str, Any]] = OrderedDict()
        captured_at = max(str(row["captured_at"]) for row in rows)
        for row in rows:
            key = (str(row["category"]), str(row["raw_course_code"]), str(row["raw_course_title"]))
            course = grouped.setdefault(
                key,
                {
                    "category": key[0],
                    "course_code": key[1],
                    "course_title": key[2],
                    "allocations": [],
                },
            )
            if row["allocation_ordinal"] is not None:
                course["allocations"].append(
                    {
                        "source_row": int(row["source_row"]),
                        "slot": str(row["raw_slot"]),
                        "venue": str(row["raw_venue"]),
                        "faculty": str(row["raw_faculty"]),
                        "course_type": str(row["raw_course_type"]),
                    }
                )
        return SemesterRows(campus, semester_label, captured_at, tuple(grouped.values()))

    def canonical_json(self, snapshot: SemesterRows) -> bytes:
        payload = {
            "schema_version": "1.0.0",
            "source": "vtop-tampermonkey",
            "campus": snapshot.campus,
            "semester_label": snapshot.semester_label,
            "captured_at": snapshot.captured_at,
            "courses": list(snapshot.courses),
        }
        return json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")

    def csv_and_sidecar(self, snapshot: SemesterRows) -> tuple[bytes, bytes]:
        stream = io.StringIO(newline="")
        writer = csv.DictWriter(stream, fieldnames=CSV_COLUMNS, lineterminator="\r\n")
        writer.writeheader()
        zero_courses: list[dict[str, str]] = []
        for course in snapshot.courses:
            if not course["allocations"]:
                zero_courses.append(
                    {key: str(course[key]) for key in ("category", "course_code", "course_title")}
                )
            for allocation in course["allocations"]:
                writer.writerow(
                    {
                        "semester_label": _spreadsheet_safe(snapshot.semester_label),
                        "category": _spreadsheet_safe(course["category"]),
                        "course_code": _spreadsheet_safe(course["course_code"]),
                        "course_title": _spreadsheet_safe(course["course_title"]),
                        "source_row": allocation["source_row"],
                        "slot": _spreadsheet_safe(allocation["slot"]),
                        "venue": _spreadsheet_safe(allocation["venue"]),
                        "faculty": _spreadsheet_safe(allocation["faculty"]),
                        "course_type": _spreadsheet_safe(allocation["course_type"]),
                    }
                )
        sidecar = {
            "schema_version": "1.0.0",
            "source": "ffcs-studio-csv-export",
            "campus": snapshot.campus,
            "semester_label": snapshot.semester_label,
            "captured_at": snapshot.captured_at,
            "zero_allocation_courses": zero_courses,
            "provenance": "Active committed local SQLite offering revisions",
            "spreadsheet_formula_neutralization": (
                "Cells starting =, +, -, or @ are prefixed with an apostrophe in CSV "
                "only; canonical JSON is unchanged."
            ),
        }
        return ("\ufeff" + stream.getvalue()).encode("utf-8"), json.dumps(
            sidecar, ensure_ascii=False, indent=2
        ).encode("utf-8")

    def xlsx(self, snapshot: SemesterRows) -> bytes:
        workbook = Workbook()
        initial_sheet = workbook.active
        if initial_sheet is not None:
            workbook.remove(initial_sheet)
        headers = [
            "Course Code",
            "Course Title",
            "Category",
            "Semester",
            "Source Row",
            "Slot",
            "Faculty",
            "Venue",
            "Component Type",
        ]
        categories = list(dict.fromkeys(str(course["category"]) for course in snapshot.courses))
        sheet_names = [*categories, "ALL_ROWS"]
        for sheet_name in sheet_names:
            sheet = workbook.create_sheet(_safe_sheet_name(sheet_name, workbook.sheetnames))
            sheet.append(headers)
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="2563EB")
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = "A1:I1"
            sheet.sheet_view.showGridLines = False
            sheet.column_dimensions["A"].width = 18
            sheet.column_dimensions["B"].width = 38
            sheet.column_dimensions["C"].width = 12
            sheet.column_dimensions["D"].width = 28
            sheet.column_dimensions["E"].width = 12
            sheet.column_dimensions["F"].width = 24
            sheet.column_dimensions["G"].width = 34
            sheet.column_dimensions["H"].width = 18
            sheet.column_dimensions["I"].width = 18
        by_name = {sheet.title: sheet for sheet in workbook.worksheets}
        category_sheets = {
            category: by_name[next(name for name in by_name if name == _safe_sheet_base(category))]
            for category in categories
        }
        all_rows = by_name["ALL_ROWS"]
        for course in snapshot.courses:
            for allocation in course["allocations"]:
                values = [
                    course["course_code"],
                    course["course_title"],
                    course["category"],
                    snapshot.semester_label,
                    allocation["source_row"],
                    allocation["slot"],
                    allocation["faculty"],
                    allocation["venue"],
                    allocation["course_type"],
                ]
                for sheet in (category_sheets[str(course["category"])], all_rows):
                    sheet.append(
                        [
                            _spreadsheet_safe(value) if isinstance(value, str) else value
                            for value in values
                        ]
                    )
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()

    def ffcsinator_zip(self, snapshot: SemesterRows) -> bytes:
        grouped: dict[str, list[dict[str, Any]]] = {}
        diagnostics: list[dict[str, Any]] = []
        for course in snapshot.courses:
            category = str(course["category"])
            grouped.setdefault(category, [])
            for allocation in course["allocations"]:
                component = str(allocation["course_type"]).upper()
                suffix = "P" if component == "ELA" else "L"
                if component not in {"TH", "ETH", "ELA", "EPJ", "SS"}:
                    diagnostics.append(
                        {
                            "course_code": course["course_code"],
                            "source_row": allocation["source_row"],
                            "warning": "Unknown component is represented on the legacy theory side",
                        }
                    )
                if component == "EPJ":
                    diagnostics.append(
                        {
                            "course_code": course["course_code"],
                            "source_row": allocation["source_row"],
                            "warning": (
                                "EPJ is lossy in FFCSinator and is represented on the L side"
                            ),
                        }
                    )
                grouped[category].append(
                    {
                        "key": f"{course['course_code']}{suffix} - {course['course_title']}",
                        "slot": allocation["slot"],
                        "venue": allocation["venue"],
                        "faculty": allocation["faculty"],
                    }
                )
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for category, rows in grouped.items():
                archive.writestr(
                    CATEGORY_FILENAMES.get(category, f"Category_{_filename_part(category)}.ts"),
                    _typescript(rows),
                )
            sidecar = {
                "schema_version": "1.0.0",
                "kind": "ffcs-studio-ffcsinator-sidecar",
                "canonical": json.loads(self.canonical_json(snapshot)),
                "diagnostics": diagnostics,
                "missing_semantics": [
                    "Legacy rows cannot distinguish TH, ETH, EPJ, and SS on the L side",
                    "Legacy rows do not carry source_row or source category inside each record",
                    "Zero-allocation courses are present only in this sidecar",
                ],
            }
            archive.writestr(
                "ffcs-studio-sidecar.json", json.dumps(sidecar, ensure_ascii=False, indent=2)
            )
        return output.getvalue()


def parse_ffcsinator_typescript(
    content: str,
    *,
    filename: str,
    campus: str,
    semester_label: str,
    captured_at: str | None = None,
) -> tuple[dict[str, Any], list[dict[str, str]]]:
    """Parse only a data-object subset. No expression or module code is executed."""

    if len(content.encode("utf-8")) > 2_000_000:
        raise InterchangeError("FFCSinator TypeScript input exceeds 2 MB")
    match = re.fullmatch(
        r"\s*(?:(?:export\s+default)|(?:export\s+)?const\s+[A-Za-z_$][\w$]*\s*=)\s*(\{.*\})\s*(?:as\s+const\s*)?;?\s*",
        content,
        flags=re.DOTALL,
    )
    if match is None:
        raise InterchangeError("Only one exported or const data object is supported")
    parsed = _TsParser(match.group(1)).parse()
    if not isinstance(parsed, dict):
        raise InterchangeError("Top-level TypeScript value must be an object")
    category = FILENAME_CATEGORIES.get(filename.casefold())
    if category is None:
        raise InterchangeError("Filename must be a documented FFCSinator category filename")
    courses: OrderedDict[tuple[str, str], dict[str, Any]] = OrderedDict()
    diagnostics: list[dict[str, str]] = []
    for key, raw_rows in parsed.items():
        if not isinstance(key, str) or not isinstance(raw_rows, list):
            raise InterchangeError("Course keys must map to arrays")
        key_match = re.fullmatch(r"(.+?)([LP])\s+-\s+(.+)", key.strip())
        if key_match is None:
            raise InterchangeError(f"Unsupported legacy course key: {key!r}")
        course_code, suffix, title = key_match.groups()
        course = courses.setdefault(
            (course_code, title),
            {
                "category": category,
                "course_code": course_code,
                "course_title": title,
                "allocations": [],
            },
        )
        for raw_row in raw_rows:
            if not isinstance(raw_row, dict) or set(raw_row) != {"slot", "venue", "faculty"}:
                raise InterchangeError("Each legacy row must contain only slot, venue, and faculty")
            if not all(isinstance(raw_row[field], str) and raw_row[field] for field in raw_row):
                raise InterchangeError("Legacy row values must be non-empty strings")
            component = (
                "UNKNOWN_LEGACY_THEORY_SIDE" if suffix == "L" else "UNKNOWN_LEGACY_PRACTICAL_SIDE"
            )
            course["allocations"].append(
                {
                    "source_row": len(course["allocations"]) + 1,
                    "slot": raw_row["slot"],
                    "venue": raw_row["venue"],
                    "faculty": raw_row["faculty"],
                    "course_type": component,
                }
            )
            diagnostics.append(
                {
                    "level": "warning",
                    "code": "legacy_component_semantics_unknown",
                    "message": (
                        f"{key}: {suffix} preserves the legacy side but not exact VTOP "
                        "component type"
                    ),
                }
            )
    timestamp = captured_at or datetime.now(UTC).isoformat()
    return (
        {
            "schema_version": "1.0.0",
            "source": "legacy-ffcsinator",
            "campus": campus,
            "semester_label": semester_label,
            "captured_at": timestamp,
            "courses": list(courses.values()),
        },
        diagnostics,
    )


class _TsParser:
    TOKEN = re.compile(
        r"\s*(?:(?P<string>'(?:\\.|[^'\\])*'|\"(?:\\.|[^\"\\])*\")|(?P<name>[A-Za-z_$][\w$]*)|(?P<punct>[{}\[\]:,]))",
        re.DOTALL,
    )

    def __init__(self, text: str) -> None:
        self.tokens: list[tuple[str, str]] = []
        cursor = 0
        while cursor < len(text):
            token = self.TOKEN.match(text, cursor)
            if token is None:
                raise InterchangeError(f"Unsupported TypeScript syntax near character {cursor}")
            kind = "string" if token.group("string") else "name" if token.group("name") else "punct"
            self.tokens.append((kind, token.group(kind)))
            cursor = token.end()
        self.index = 0

    def parse(self) -> Any:
        value = self.value()
        if self.index != len(self.tokens):
            raise InterchangeError("Unexpected tokens after data object")
        return value

    def value(self) -> Any:
        kind, value = self.peek()
        if value == "{":
            return self.object()
        if value == "[":
            return self.array()
        if kind == "string":
            self.index += 1
            parsed = ast.literal_eval(value)
            if not isinstance(parsed, str):
                raise InterchangeError("Only string literals are supported")
            return parsed
        raise InterchangeError("Only objects, arrays, and string literals are supported")

    def object(self) -> dict[str, Any]:
        self.expect("{")
        result: dict[str, Any] = {}
        while self.peek()[1] != "}":
            kind, raw_key = self.peek()
            if kind == "string":
                key = self.value()
            elif kind == "name":
                self.index += 1
                key = raw_key
            else:
                raise InterchangeError("Object keys must be strings or identifiers")
            self.expect(":")
            result[key] = self.value()
            if self.peek()[1] == ",":
                self.index += 1
                if self.peek()[1] == "}":
                    break
            elif self.peek()[1] != "}":
                raise InterchangeError("Expected comma in object")
        self.expect("}")
        return result

    def array(self) -> list[Any]:
        self.expect("[")
        result: list[Any] = []
        while self.peek()[1] != "]":
            result.append(self.value())
            if self.peek()[1] == ",":
                self.index += 1
                if self.peek()[1] == "]":
                    break
            elif self.peek()[1] != "]":
                raise InterchangeError("Expected comma in array")
        self.expect("]")
        return result

    def peek(self) -> tuple[str, str]:
        if self.index >= len(self.tokens):
            raise InterchangeError("Unexpected end of TypeScript data")
        return self.tokens[self.index]

    def expect(self, value: str) -> None:
        if self.peek()[1] != value:
            raise InterchangeError(f"Expected {value!r}")
        self.index += 1


def inspect_xlsx_rows(content: bytes) -> list[tuple[Any, ...]]:
    """Test helper that reads the exported searchable sheet without formulas."""

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    return [tuple(row) for row in workbook["ALL_ROWS"].iter_rows(values_only=True)]


def _spreadsheet_safe(value: object) -> str:
    text = str(value)
    return "'" + text if text.lstrip().startswith(("=", "+", "-", "@")) else text


def _safe_sheet_base(value: str) -> str:
    cleaned = re.sub(r"[\\/*?:\[\]]", "_", value).strip(" '") or "CATEGORY"
    return cleaned[:31]


def _safe_sheet_name(value: str, existing: list[str]) -> str:
    base = _safe_sheet_base(value)
    candidate = base
    number = 2
    while candidate in existing:
        suffix = f"_{number}"
        candidate = base[: 31 - len(suffix)] + suffix
        number += 1
    return candidate


def _filename_part(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9_-]+", "_", value)[:60] or "Unknown"


def _typescript(rows: list[dict[str, Any]]) -> str:
    grouped: OrderedDict[str, list[dict[str, str]]] = OrderedDict()
    for row in rows:
        grouped.setdefault(str(row["key"]), []).append(
            {field: str(row[field]) for field in ("slot", "venue", "faculty")}
        )
    body = json.dumps(grouped, ensure_ascii=False, indent=2)
    return (
        "// Generated by FFCS Studio 0.7; data only, no executable input.\nexport default "
        + body
        + ";\n"
    )
